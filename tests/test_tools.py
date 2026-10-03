import json
import zipfile

import pytest
import requests

from data.sample.companies import COMPANIES, all_pages, company_pages
from tools import search as S
from tools.docx_writer import docxjs_available, read_docx_skill, write_proposal
from tools.extract import find_links, html_to_text, key_pages
from tools.fetch import fetch_site
from tools.field_classifier import extract_fields
from tools.http import FetchError, HttpClient, RobotsDisallowed
from tools.research import OfflineResearchProvider, WikipediaProvider, research_field
from tools.telegram import TelegramClient, chunk_text

HTTP_CFG = {"user_agent": "TestBot/1.0", "timeout_s": 5, "retries": 2, "min_delay_s": 0.0, "cache": False}


class FakeResp:
    def __init__(self, status=200, text="", js=None, headers=None):
        self.status_code, self.text, self._js = status, text, js
        self.headers = headers or {}
        self.encoding = "utf-8"
        self.url = "https://x.example/"

    def iter_content(self, n):
        yield self.text.encode()

    def json(self):
        if self._js is None:
            raise ValueError
        return self._js


class FakeSession:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []
        self.headers = {}

    def _next(self, *a, **kw):
        self.calls.append((a, kw))
        r = self.responses.pop(0)
        if isinstance(r, Exception):
            raise r
        return r

    get = post = _next


# ------------------------------------------------------------------ http
def test_robots_respected():
    http = HttpClient(HTTP_CFG, offline=all_pages())
    c = COMPANIES[0]
    assert http.allowed(f"https://{c['domain']}/about")
    assert not http.allowed(f"https://{c['domain']}/admin/panel")
    with pytest.raises(RobotsDisallowed):
        http.get(f"https://{c['domain']}/admin/panel")


def test_retry_then_success_and_rate_limit():
    sleeps = []
    sess = FakeSession([FakeResp(503), FakeResp(200, "ok")])
    http = HttpClient(dict(HTTP_CFG, min_delay_s=0.0, respect_robots=False), session=sess, sleep=sleeps.append)
    assert http.get("https://x.example/a").text == "ok"
    assert len(sess.calls) == 2 and sleeps  # backed off once


def test_network_down_raises_after_retries():
    sess = FakeSession([requests.ConnectionError("down")] * 3)
    http = HttpClient(dict(HTTP_CFG, respect_robots=False), session=sess, sleep=lambda s: None)
    with pytest.raises(requests.ConnectionError):
        http.get("https://x.example/")
    assert len(sess.calls) == 3


def test_per_host_politeness_delay():
    sleeps = []
    sess = FakeSession([FakeResp(200, "a"), FakeResp(200, "b")])
    http = HttpClient(dict(HTTP_CFG, min_delay_s=2.0, respect_robots=False), session=sess, sleep=sleeps.append)
    http.get("https://x.example/1")
    http.get("https://x.example/2")
    assert any(1.0 < s <= 2.0 for s in sleeps)


def test_cache(tmp_path):
    sess = FakeSession([FakeResp(200, "cached body")])
    http = HttpClient(dict(HTTP_CFG, respect_robots=False, cache=True), cache_dir=tmp_path, session=sess)
    assert http.get("https://x.example/c").text == "cached body"
    r = http.get("https://x.example/c")
    assert r.from_cache and len(sess.calls) == 1


# ------------------------------------------------------------------ extract / fetch / classify
def test_html_to_text_removes_boilerplate():
    html = company_pages(COMPANIES[0])[f"https://{COMPANIES[0]['domain']}/about"]
    text = html_to_text(html)
    assert "Hamburg" in text and "cookies" not in text.lower() and "All rights reserved" not in text
    assert "var tracking" not in text


def test_bad_html_does_not_crash():
    assert isinstance(html_to_text("<html><body><p>unclosed <div><<<>>>\x00\xff garbage"), str)
    assert html_to_text("") == ""


def test_key_pages_and_fetch_site():
    c = COMPANIES[2]
    http = HttpClient(HTTP_CFG, offline=all_pages())
    home = company_pages(c)[f"https://{c['domain']}/"]
    links = find_links(home, f"https://{c['domain']}/")
    kp = key_pages(links)
    assert any("about" in u for u in kp) and any("services" in u for u in kp) and any("products" in u for u in kp)
    site = fetch_site(http, f"https://{c['domain']}/")
    assert len(site["pages"]) >= 4 and "Toronto" in " ".join(site["pages"].values())


def test_fetch_site_homepage_missing():
    http = HttpClient(HTTP_CFG, offline={})
    with pytest.raises(FetchError):
        fetch_site(http, "https://nowhere.example/")


@pytest.mark.parametrize("c", COMPANIES, ids=[c["domain"] for c in COMPANIES])
def test_field_classifier_on_fixtures(c):
    text = "\n".join(html_to_text(h) for u, h in company_pages(c).items() if not u.endswith("robots.txt"))
    info = extract_fields(text)
    assert info["field"] == c["field"]
    assert info["country"] == c["country"]
    assert info["name"] == c["name"]
    assert info["size"] == c["employees"]
    assert info["services"]


# ------------------------------------------------------------------ search
DDG_HTML = """<html><body>
<a class="result__a" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fwww.acme-solar.example%2Fabout&rut=x">Acme Solar</a>
<a class="result__a" href="https://www.linkedin.com/company/acme">Acme on LinkedIn</a>
<a class="result__a" href="https://acme-solar.example/contact">Acme contact</a>
<a class="result__a" href="https://beta-wind.example/">Beta Wind</a></body></html>"""


def test_duckduckgo_parse_filters_and_dedups():
    http = HttpClient(HTTP_CFG, offline={S.DuckDuckGoProvider.ENDPOINT: DDG_HTML})
    res = S.DuckDuckGoProvider(http).search("Germany", "renewable energy", 10)
    assert [r["domain"] for r in res] == ["acme-solar.example", "beta-wind.example"]
    assert res[0]["url"] == "https://www.acme-solar.example/"


def test_seedfile_and_offline_providers(tmp_path):
    p = tmp_path / "seed.json"
    p.write_text(json.dumps([{"name": "A", "url": "https://a.example/x", "country": "Germany", "industry": "logistics"},
                             {"name": "B", "url": "https://b.example", "country": "France", "industry": "logistics"}]))
    assert [r["domain"] for r in S.SeedFileProvider(str(p)).search("Germany", "logistics", 5)] == ["a.example"]
    off = S.OfflineProvider().search("Germany", "", 10)
    assert {r["domain"] for r in off} == {"nordwind-solar.example", "rheinland-freight.example"}


def test_gather_dedup_and_failure_isolation():
    class Flaky:
        name = "flaky"

        def search(self, country, industry, limit):
            if country == "Canada":
                raise requests.ConnectionError("network down")
            return [{"domain": "same.example", "url": "https://same.example/"},
                    {"domain": f"{country}.example", "url": "u"}]
    res = S.gather_companies(Flaky(), ["Germany", "Canada", "France"], ["x"], 10, seen_domains={"France.example"})
    assert [r["domain"] for r in res] == ["same.example", "Germany.example"]


# ------------------------------------------------------------------ research
def test_offline_research_has_urls():
    ev = research_field(OfflineResearchProvider(), "logistics", [], 6)
    assert ev and all(e["url"].startswith("https://") and e["id"].startswith("E") for e in ev)


def test_wikipedia_provider_parses_api():
    def fake(url, params):
        if params.get("list") == "search":
            return json.dumps({"query": {"search": [{"title": "Solar power"}]}})
        return json.dumps({"query": {"pages": {"1": {"extract": "Solar power is the conversion of sunlight. " * 20}}}})
    http = HttpClient(HTTP_CFG, offline=fake)
    docs = WikipediaProvider(http).research("renewable energy", ["solar"], 3)
    assert docs[0]["url"] == "https://en.wikipedia.org/wiki/Solar_power" and "sunlight" in docs[0]["text"]


# ------------------------------------------------------------------ docx
SECTIONS = [
    {"title": "Cover", "paragraphs": ["Project Proposal for Test Co", "Field: software"], "bullets": [], "table": None},
    {"title": "Executive Summary", "paragraphs": ["Summary sentence. [S1]"], "bullets": ["one", "two"], "table": None},
    {"title": "Timeline and Milestones", "paragraphs": [], "bullets": [],
     "table": {"headers": ["Phase", "Weeks", "Milestone"], "rows": [["Discovery", "1–3", "Baseline"]]}},
    {"title": "Sources", "paragraphs": ["[S1] https://kb.example/x"], "bullets": [], "table": None},
]


@pytest.mark.parametrize("backend", ["python-docx", "docxjs"])
def test_docx_structure(tmp_path, backend):
    if backend == "docxjs" and not docxjs_available():
        pytest.skip("node/docx not available")
    info = write_proposal(SECTIONS, tmp_path / "p.docx", "Proposal — Test Co", backend=backend)
    assert info["backend"] == backend
    with zipfile.ZipFile(tmp_path / "p.docx") as z:
        doc = z.read("word/document.xml").decode()
        footers = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))
    assert "TOC" in doc                         # table of contents field
    assert "PAGE" in footers                    # page numbers
    assert "Executive Summary" in doc and "Discovery" in doc and "<w:tbl>" in doc
    import docx
    d = docx.Document(str(tmp_path / "p.docx"))
    heads = [p.text for p in d.paragraphs if p.style is not None and p.style.name.startswith("Heading")]
    assert "Executive Summary" in heads and "Sources" in heads
    assert len(d.tables) == 1


def test_docx_skill_read():
    info = read_docx_skill("/mnt/skills/public/docx/SKILL.md")
    if info["available"]:
        assert any("Tables need dual widths" in r for r in info["rules"])
    assert read_docx_skill("/nonexistent/SKILL.md")["available"] is False


# ------------------------------------------------------------------ telegram
TG_CFG = {"telegram": {"enabled": True, "chunk_chars": 50, "retries": 2, "timeout_s": 5}}


def test_chunk_text():
    text = "\n".join(f"line {i} " + "x" * 20 for i in range(20))
    chunks = chunk_text(text, 100)
    assert all(len(c) <= 100 for c in chunks) and "".join(chunks) == text
    assert all(len(c) <= 10 for c in chunk_text("y" * 35, 10))


def test_telegram_chunked_delivery_confirmation():
    sess = FakeSession([FakeResp(200, js={"ok": True, "result": {"message_id": i}}) for i in range(1, 10)])
    tg = TelegramClient(TG_CFG, token="T", chat_id="C", session=sess, sleep=lambda s: None)
    rec = tg.send_message("a" * 120)
    assert len(rec) == 3 and all(r["delivered"] for r in rec) and [r["message_id"] for r in rec] == [1, 2, 3]


def test_telegram_retry_after_429(tmp_path):
    sleeps = []
    sess = FakeSession([FakeResp(429, js={"ok": False, "description": "Too Many Requests", "parameters": {"retry_after": 3}}),
                        FakeResp(200, js={"ok": True, "result": {"message_id": 7}})])
    tg = TelegramClient(TG_CFG, token="T", chat_id="C", session=sess, sleep=sleeps.append)
    f = tmp_path / "p.docx"
    f.write_bytes(b"x")
    rec = tg.send_document(f, "caption")
    assert rec["delivered"] and rec["message_id"] == 7 and sleeps == [3]


def test_telegram_error_and_not_configured(tmp_logs):
    sess = FakeSession([FakeResp(400, js={"ok": False, "description": "Bad Request: chat not found"})] * 3)
    tg = TelegramClient(TG_CFG, token="T", chat_id="C", session=sess, sleep=lambda s: None)
    rec = tg.send_message("hello")
    assert not rec[0]["delivered"] and "chat not found" in rec[0]["error"]
    assert "chat not found" in (tmp_logs / "errors.log").read_text()
    off = TelegramClient(TG_CFG, token=None, chat_id=None)
    off.token = off.chat_id = None
    assert off.send_message("x")[0]["error"] == "not configured"
