"""Write a polished proposal .docx: cover page, table of contents, headings, tables, bullets, page numbers.

Backends:
  docxjs       — Node + docx-js (`tools/docx_render.js`), following the Word/docx skill instructions
                 (`paths.docx_skill`, e.g. /mnt/skills/public/docx/SKILL.md) when available.
  python-docx  — pure-Python fallback with the same structure.
"""
from __future__ import annotations

import functools
import json
import os
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from core.config import ROOT
from core.logging import get_logger

log = get_logger("tools")


def read_docx_skill(path: str | Path) -> dict:
    """To-do item 5: fetch and read the Word/docx skill; return the rules we apply."""
    p = Path(path)
    if not p.exists():
        return {"available": False, "path": str(p), "backend": "python-docx", "rules": []}
    text = p.read_text(encoding="utf-8")
    rules = [ln.strip("- ").strip() for ln in text.splitlines() if ln.strip().startswith("- **")]
    return {"available": True, "path": str(p), "backend": "docxjs", "rules": rules, "chars": len(text)}


@functools.lru_cache(maxsize=1)
def _node_env() -> dict | None:
    node = shutil.which("node")
    if not node:
        return None
    env = dict(os.environ)
    paths = [str(ROOT / "node_modules")]
    try:
        paths.append(subprocess.run(["npm", "root", "-g"], capture_output=True, text=True, timeout=20).stdout.strip())
    except (OSError, subprocess.SubprocessError):
        pass
    env["NODE_PATH"] = os.pathsep.join(p for p in paths if p)
    ok = subprocess.run([node, "-e", "require('docx')"], env=env, capture_output=True, timeout=30).returncode == 0
    return env if ok else None


def docxjs_available() -> bool:
    return _node_env() is not None


def _write_docxjs(spec: dict, out: Path) -> None:
    env = _node_env()
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as f:
        json.dump(spec, f, ensure_ascii=False)
        spec_path = f.name
    try:
        r = subprocess.run(["node", str(ROOT / "tools" / "docx_render.js"), spec_path, str(out)], env=env,
                           capture_output=True, text=True, timeout=120)
        if r.returncode != 0:
            raise RuntimeError(f"docx-js failed: {r.stderr[-800:]}")
    finally:
        os.unlink(spec_path)


def _add_field(paragraph, instr: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _write_python_docx(spec: dict, out: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    _add_field(fp, "PAGE")
    fp.add_run(" of ")
    _add_field(fp, "NUMPAGES")
    sec.header.paragraphs[0].text = spec["title"]
    cover = next((s for s in spec["sections"] if s["title"] == "Cover"), {"paragraphs": [spec["title"]]})
    for _ in range(8):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(cover["paragraphs"][0])
    r.bold, r.font.size, r.font.color.rgb = True, Pt(26), RGBColor(0x1F, 0x4E, 0x79)
    for line in cover["paragraphs"][1:]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Table of Contents", level=1)
    _add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for s in spec["sections"]:
        if s["title"] == "Cover":
            continue
        doc.add_heading(s["title"], level=1)
        for p in s.get("paragraphs", []):
            doc.add_paragraph(p)
        for b in s.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")
        tbl = s.get("table")
        if tbl and tbl.get("rows"):
            table = doc.add_table(rows=1, cols=len(tbl["headers"]))
            table.style = "Light Grid Accent 1"
            for i, h in enumerate(tbl["headers"]):
                table.rows[0].cells[i].text = h
            for row in tbl["rows"]:
                cells = table.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
            doc.add_paragraph()
    _ = WD_SECTION
    doc.save(str(out))


def write_proposal(sections: list[dict], out: Path, title: str, backend: str = "auto", page_size: str = "A4",
                   company: str | None = None) -> dict:
    """Render sections to `out`. Returns {"path", "backend", "bytes"}."""
    out.parent.mkdir(parents=True, exist_ok=True)
    spec = {"title": title, "page_size": page_size, "sections": sections}
    t0 = time.time()
    chosen = backend
    if backend == "auto":
        chosen = "docxjs" if docxjs_available() else "python-docx"
    try:
        if chosen == "docxjs":
            _write_docxjs(spec, out)
        else:
            _write_python_docx(spec, out)
    except Exception as e:
        if chosen == "docxjs":
            log.event("docx_writer", "warn", tool="docx", company=company, details=f"docx-js failed, falling back: {e}")
            chosen = "python-docx"
            _write_python_docx(spec, out)
        else:
            raise
    info = {"path": str(out), "backend": chosen, "bytes": out.stat().st_size}
    log.event("docx_writer", "ok", tool="docx", company=company, duration=time.time() - t0, details=info)
    return info
